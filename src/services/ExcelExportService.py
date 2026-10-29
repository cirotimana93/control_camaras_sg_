import io
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExcelExportService:

    @staticmethod
    def getMonthlyLegendDataFrame() -> pd.DataFrame:
        legendData = [
            {
                "Concepto / Columna": "GRABANDO",
                "Tipo": "Estado Actual",
                "Criticidad": "NORMAL",
                "Significado Tecnico": "La camara se encuentra con enlace de red y registrando video continuo en el disco duro del DVR al momento del reporte.",
                "Accion Recomendada": "Ninguna accion requerida. Operatividad normal."
            },
            {
                "Concepto / Columna": "ONLINE NO GRABA",
                "Tipo": "Estado Actual",
                "Criticidad": "ALTA",
                "Significado Tecnico": "La camara responde en red pero la grabacion en disco esta detenida. Causa probable: disco duro lleno/dañado o programacion de grabacion desactivada.",
                "Accion Recomendada": "Ingresar a la interfaz web del DVR, revisar salud del disco HDD y verificar que el canal tenga activo el horario de grabacion continua."
            },
            {
                "Concepto / Columna": "OFFLINE",
                "Tipo": "Estado Actual",
                "Criticidad": "MEDIA",
                "Significado Tecnico": "La camara no entrega senal de video hacia el DVR. Causa probable: transformador/fuente 12V apagada o cable de video BNC/UTP desconectado.",
                "Accion Recomendada": "Revisar la fuente de poder de la camara y conexionado fisico del cable en tienda."
            },
            {
                "Concepto / Columna": "SIN REPORTE",
                "Tipo": "Estado Actual",
                "Criticidad": "ALTA",
                "Significado Tecnico": "El DVR completo no ha transmitido telemetria al servidor central en las ultimas 48 horas.",
                "Accion Recomendada": "Verificar encendido del DVR y conexion del cable de red/internet de la tienda."
            },
            {
                "Concepto / Columna": "dias sin grabacion acumulado",
                "Tipo": "Metrica Historica",
                "Criticidad": "INFORMATIVA",
                "Significado Tecnico": "Total de dias calendario en los que la camara tuvo valor '0' (sin grabacion) en el calendario de los ultimos 30 dias.",
                "Accion Recomendada": "Priorizar camaras con mas de 5 dias acumulados sin grabacion para auditoria."
            },
            {
                "Concepto / Columna": "dias en los que dejo de grabar",
                "Tipo": "Metrica Historica",
                "Criticidad": "INFORMATIVA",
                "Significado Tecnico": "Listado de fechas especificas (dia/mes) o rangos de dias donde se identifico perdida o interrupcion de grabacion.",
                "Accion Recomendada": "Utilizar como sustento para identificar eventos de corte o fechas de inicio de falla."
            }
        ]
        return pd.DataFrame(legendData)

    @staticmethod
    def getDailyLegendDataFrame() -> pd.DataFrame:
        legendData = [
            {
                "Estado Actual": "OFFLINE / SIN REPORTE (>24H)",
                "Criticidad": "ALTA",
                "Significado Tecnico": "El equipo DVR o agente de tienda no transmite telemetria hace mas de 24 horas. Causa probable: corte de energia prolongado, desconexion de red o corte de internet.",
                "Accion Recomendada": "Contactar con tienda o soporte de campo para verificar encendido del DVR, luces del puerto Ethernet y conexion de enlace de internet."
            },
            {
                "Estado Actual": "SIN REPORTE (>30 MIN)",
                "Criticidad": "MEDIA",
                "Significado Tecnico": "El DVR dejo de enviar reporte hace mas de 30 minutos (reciente). Causa probable: microcorte de enlace, reinicio de equipo o corte temporal de energia.",
                "Accion Recomendada": "Monitorear reconexion automatica. Si el tiempo supera 2 horas, reportar a soporte tecnico."
            },
            {
                "Estado Actual": "ONLINE NO GRABA",
                "Criticidad": "ALTA",
                "Significado Tecnico": "El DVR y la camara estan en linea en la red, pero el canal especifico no registra video en disco duro.",
                "Accion Recomendada": "Ingresar a la interfaz web del DVR, revisar salud del disco duro HDD y activar el horario de grabacion continua en dicho canal."
            },
            {
                "Estado Actual": "CAMARA OFFLINE",
                "Criticidad": "MEDIA",
                "Significado Tecnico": "El DVR responde pero la camara individual no entrega senal de video. Causa probable: fuente 12V desconectada/quemada o cable BNC/UTP suelto.",
                "Accion Recomendada": "Revisar transformador de corriente de la camara y conexionado fisico del canal."
            },
            {
                "Estado Actual": "NO GRABA NINGUN CANAL",
                "Criticidad": "CRITICA",
                "Significado Tecnico": "El DVR esta encendido pero ninguno de sus canales configurados esta grabando. Causa probable: falla total de disco duro o disco no formateado.",
                "Accion Recomendada": "Diagnostico urgente de almacenamiento o reemplazo de disco rigido."
            },
            {
                "Estado Actual": "GRABACION PARCIAL (X/Y)",
                "Criticidad": "MEDIA",
                "Significado Tecnico": "X canales graban correctamente de un total de Y canales configurados en la tienda.",
                "Accion Recomendada": "Inspeccionar unicamente los canales con estado 'ONLINE NO GRABA' o 'CAMARA OFFLINE'."
            },
            {
                "Estado Actual": "OBSERVADO",
                "Criticidad": "BAJA",
                "Significado Tecnico": "Canal con cortes intermitentes detectados en las ultimas horas.",
                "Accion Recomendada": "Monitorear comportamiento durante el dia."
            },
            {
                "Estado Actual": "GRABANDO",
                "Criticidad": "NORMAL",
                "Significado Tecnico": "Operatividad normal. La camara y el DVR estan sincronizados y almacenando video continuo en disco.",
                "Accion Recomendada": "Ninguna accion requerida."
            }
        ]
        return pd.DataFrame(legendData)

    @staticmethod
    def styleWorksheet(worksheet, headerColorHex: str = "1F4E79"):
        headerFill = PatternFill(start_color=headerColorHex, end_color=headerColorHex, fill_type="solid")
        headerFont = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regularFont = Font(name="Calibri", size=10)
        thinBorder = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )

        for colIdx in range(1, worksheet.max_column + 1):
            cell = worksheet.cell(row=1, column=colIdx)
            cell.fill = headerFill
            cell.font = headerFont
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        worksheet.row_dimensions[1].height = 28

        for rowIdx in range(2, worksheet.max_row + 1):
            worksheet.row_dimensions[rowIdx].height = 22
            for colIdx in range(1, worksheet.max_column + 1):
                cell = worksheet.cell(row=rowIdx, column=colIdx)
                cell.font = regularFont
                cell.border = thinBorder
                cell.alignment = Alignment(vertical="center", wrap_text=True)

        # auto-ajuste inteligente de ancho de columnas
        for col in worksheet.columns:
            maxLen = 0
            colLetter = get_column_letter(col[0].column)
            for cell in col:
                valStr = str(cell.value or "")
                if len(valStr) > maxLen:
                    maxLen = len(valStr)
            worksheet.column_dimensions[colLetter].width = min(max(maxLen + 4, 14), 50)

    @staticmethod
    def createMonthlyExcel(cameraDf: pd.DataFrame, summaryDf: pd.DataFrame) -> bytes:
        outputBuffer = io.BytesIO()
        legendDf = ExcelExportService.getMonthlyLegendDataFrame()
        
        # filtrar camaras observadas
        observedFilter = (cameraDf["dias sin grabacion acumulado"] > 0) | (cameraDf["estado actual"] != "GRABANDO")
        observedDf = cameraDf[observedFilter].copy()
        
        with pd.ExcelWriter(outputBuffer, engine="openpyxl") as writer:
            summaryDf.to_excel(writer, sheet_name="Resumen por DVR", index=False)
            observedDf.to_excel(writer, sheet_name="Camaras Observadas", index=False)
            cameraDf.to_excel(writer, sheet_name="Todas las Camaras", index=False)
            legendDf.to_excel(writer, sheet_name="Leyenda Mensual", index=False)
            
            ExcelExportService.styleWorksheet(writer.sheets["Resumen por DVR"], "1F4E79")
            ExcelExportService.styleWorksheet(writer.sheets["Camaras Observadas"], "C00000")
            ExcelExportService.styleWorksheet(writer.sheets["Todas las Camaras"], "1F4E79")
            ExcelExportService.styleWorksheet(writer.sheets["Leyenda Mensual"], "333F48")

        outputBuffer.seek(0)
        return outputBuffer.getvalue()

    @staticmethod
    def createDailyAlertsExcel(cameraDf: pd.DataFrame, summaryDf: pd.DataFrame) -> bytes:
        outputBuffer = io.BytesIO()
        legendDf = ExcelExportService.getDailyLegendDataFrame()
        
        with pd.ExcelWriter(outputBuffer, engine="openpyxl") as writer:
            cameraDf.to_excel(writer, sheet_name="Alertas por Camara", index=False)
            summaryDf.to_excel(writer, sheet_name="Resumen por DVR", index=False)
            legendDf.to_excel(writer, sheet_name="Leyenda Alertas Diarias", index=False)
            
            ExcelExportService.styleWorksheet(writer.sheets["Alertas por Camara"], "C00000")
            ExcelExportService.styleWorksheet(writer.sheets["Resumen por DVR"], "1F4E79")
            ExcelExportService.styleWorksheet(writer.sheets["Leyenda Alertas Diarias"], "333F48")

        outputBuffer.seek(0)
        return outputBuffer.getvalue()
